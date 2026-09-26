"""Procurement helpers: local parsing/comparison with small, explicit Gemini calls."""
from __future__ import annotations

import io
import json
import mimetypes
import os
import re
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from difflib import SequenceMatcher
from email import policy
from email.parser import BytesParser
from html import unescape
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document

BASE_DIR = Path(__file__).resolve().parent
DEFAULT_RFX_PATH = BASE_DIR / "RFx.docx"
MODEL = "gemini-3.5-flash-lite"
MAX_SOURCE_TEXT_CHARS = 50000
MAX_OUTPUT_TOKENS = 65536
VENDOR_RESPONSE_SCHEMA = {
    "type": "object",
    "properties": {
        "vendors": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "file_name": {"type": "string"},
                    "vendor": {"type": "string"},
                    "source_status": {"type": "string", "enum": ["readable", "partially_readable", "unreadable"]},
                    "source_notes": {"type": "string"},
                    "items": {"type": "array", "items": {"type": "array", "items": {"type": "string"}}},
                    "questionnaire": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {"question": {"type": "string"}, "answer": {"type": "string"}},
                            "required": ["question", "answer"],
                        },
                    },
                    "commercial_terms": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {"term": {"type": "string"}, "response": {"type": "string"}},
                            "required": ["term", "response"],
                        },
                    },
                },
                "required": ["file_name", "vendor", "source_status", "source_notes", "items", "questionnaire", "commercial_terms"],
            },
        },
    },
    "required": ["vendors"],
}


def _client():
    api_key = os.getenv("GEMINI_API_KEY") or os.getenv("GOOGLE_API_KEY")
    if not api_key:
        try:
            import streamlit as st
            api_key = st.secrets.get("GEMINI_API_KEY") or st.secrets.get("GOOGLE_API_KEY")
        except Exception:
            api_key = None
    if not api_key:
        raise RuntimeError("Set GEMINI_API_KEY (or GOOGLE_API_KEY) to use Gemini features.")
    try:
        from google import genai
    except ImportError as exc:
        raise RuntimeError("Gemini SDK is missing. Install dependencies from requirements.txt.") from exc
    return genai.Client(api_key=api_key)


def _json(text: str) -> dict:
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", (text or "").strip(), flags=re.I)
    try:
        result = json.loads(text)
    except json.JSONDecodeError as initial_error:
        start, end = text.find("{"), text.rfind("}")
        if start < 0 or end <= start:
            raise ValueError(f"Gemini JSON is incomplete: {initial_error.msg} at character {initial_error.pos}.") from initial_error
        result = json.loads(text[start:end + 1])
    if not isinstance(result, dict):
        raise ValueError("Expected a JSON object from Gemini.")
    return result


def _gemini_finish_status(response: Any) -> str:
    """Return safe response status metadata without exposing supplier content."""
    statuses = []
    prompt_feedback = getattr(response, "prompt_feedback", None)
    block_reason = getattr(prompt_feedback, "block_reason", None)
    if block_reason:
        statuses.append(f"prompt block reason: {getattr(block_reason, 'name', block_reason)}")
    candidates = getattr(response, "candidates", None) or []
    if candidates:
        finish_reason = getattr(candidates[0], "finish_reason", None)
        if finish_reason:
            statuses.append(f"finish reason: {getattr(finish_reason, 'name', finish_reason)}")
    return "; ".join(statuses) or "finish status not reported"


def _text(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return str(value).strip()


def _number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and not pd.isna(value):
        return float(value)
    match = re.search(r"-?\d[\d,]*(?:\.\d+)?", str(value).replace(",", ""))
    return float(match.group(0)) if match else None


def load_rfx(path: str | Path | bytes | bytearray = DEFAULT_RFX_PATH) -> dict:
    """Load RFx scope, line items, questionnaire, and terms from a DOCX path or upload."""
    doc = Document(io.BytesIO(path)) if isinstance(path, (bytes, bytearray)) else Document(path)
    paras = [_text(p.text) for p in doc.paragraphs if _text(p.text)]
    scope = next((p.split(":", 1)[1].strip() for p in paras if p.lower().startswith("scope:")), "")
    items = []
    for table in doc.tables:
        for row in table.rows[1:]:
            cells = [_text(c.text).replace("�", "x") for c in row.cells]
            if len(cells) >= 5 and cells[0].strip("# ").isdigit():
                items.append({"item_id": int(cells[0].strip("# ")), "item": cells[1],
                              "specification": cells[2], "quantity": _number(cells[3]), "unit": cells[4]})
    questions, in_questions = [], False
    for paragraph in paras:
        label = paragraph.lower().rstrip(":")
        if label == "questionnaire":
            in_questions = True
            continue
        if label == "commercial terms":
            in_questions = False
            continue
        if in_questions and paragraph.endswith("?"):
            questions.append(paragraph)
    terms = {}
    term_labels = {
        "quote validity": "quote_validity", "prices": "currency", "taxes": "taxes",
        "delivery": "delivery_location", "payment": "payment_terms", "freight": "freight",
        "delivery requirement": "delivery_requirement",
    }
    for paragraph in paras:
        if ":" not in paragraph:
            continue
        label, value = paragraph.split(":", 1)
        key = term_labels.get(label.lower().strip())
        if key and value.strip():
            terms[key] = value.strip()
    return {"scope": scope, "items": items, "questionnaire": questions, "commercial_terms": terms}


def empty_rfx() -> dict:
    """Return a neutral starting state; procurement facts must come from the buyer."""
    return {"scope": "", "items": [], "questionnaire": [], "commercial_terms": {}}


def _currency_code(value: Any) -> str:
    text = _text(value).upper()
    aliases = {"₹": "INR", "RS": "INR", "RUPEE": "INR", "RUPEES": "INR",
               "US DOLLAR": "USD", "US DOLLARS": "USD"}
    if text in aliases:
        return aliases[text]
    if "₹" in text or re.search(r"\bRUPEES?\b", text):
        return "INR"
    if re.search(r"\bUS DOLLARS?\b", text):
        return "USD"
    ignored = {"ALL", "THE", "FOR", "AND", "ARE", "PER", "ANY", "NOT", "BUY", "NET", "GST", "TAX"}
    return next((code for code in re.findall(r"\b[A-Z]{3}\b", text) if code not in ignored), "")


def infer_rfx_currency(rfx: dict | None) -> str:
    """Read the buyer's comparison currency from the RFx commercial terms."""
    terms = (rfx or {}).get("commercial_terms") or {}
    for key, value in terms.items():
        normalized_key = _key(_text(key))
        if normalized_key in {"currency", "pricecurrency", "comparisoncurrency", "basecurrency", "prices"}:
            code = _currency_code(value)
            if code:
                return code
    return ""


def _fetch_fx_rates(base_currency: str, source_currencies: list[str]) -> tuple[dict[str, float], dict[str, str], str | None]:
    """Fetch all required latest source-to-base rates with one no-key FX request."""
    base = _currency_code(base_currency)
    sources = sorted({code for value in source_currencies if (code := _currency_code(value)) and code != base})
    if not base or not sources:
        return {}, {}, None
    query = urllib.parse.urlencode({"base": base, "quotes": ",".join(sources)})
    url = f"https://api.frankfurter.dev/v2/rates?{query}"
    try:
        request = urllib.request.Request(url, headers={"Accept": "application/json", "User-Agent": "Aerchain-Procurement-Copilot/1.0"})
        with urllib.request.urlopen(request, timeout=8) as response:
            rows = json.load(response)
        if not isinstance(rows, list):
            raise ValueError("FX provider returned an unexpected response format.")
        # API rate is units of source quote currency per one target-base unit;
        # invert it to convert vendor prices into the comparison currency.
        rates, dates = {}, {}
        for row in rows:
            if not isinstance(row, dict):
                continue
            quote = _currency_code(row.get("quote"))
            rate = _number(row.get("rate"))
            if quote in sources and rate and rate > 0:
                rates[quote] = 1 / rate
                dates[quote] = _text(row.get("date"))
        missing = sorted(set(sources) - set(rates))
        note = f"No rate returned for {', '.join(missing)}." if missing else None
        return rates, dates, note
    except (OSError, urllib.error.URLError, TimeoutError, ValueError, json.JSONDecodeError) as exc:
        return {}, {}, f"Could not retrieve exchange rates: {exc}"


def _source(source) -> tuple[str, bytes]:
    name = getattr(source, "name", None) or Path(source).name
    if hasattr(source, "getvalue"):
        data = source.getvalue()
    elif hasattr(source, "read"):
        data = source.read()
        if hasattr(source, "seek"):
            source.seek(0)
    else:
        data = Path(source).read_bytes()
    return Path(name).name, data


def _docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    lines = [_text(p.text) for p in doc.paragraphs if _text(p.text)]
    for table in doc.tables:
        lines.extend(" | ".join(_text(cell.text) for cell in row.cells) for row in table.rows)
    return "\n".join(lines)


def _pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        from PyPDF2 import PdfReader
    return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages).strip()


def _map_headers(columns) -> dict[str, str]:
    aliases = {
        "item": ["item", "item name", "product", "description", "material"],
        "specification": ["specification", "spec", "dimensions", "details"],
        "quantity": ["quantity", "quoted quantity", "qty"],
        "unit": ["unit", "uom", "unit of measure"],
        "unit_price": ["unit price", "price", "rate", "quoted price", "price per unit"],
        "currency": ["currency", "ccy"], "gst": ["gst", "gst %", "tax", "tax %"],
        "delivery_days": ["delivery days", "delivery", "lead time"],
        "price_basis": ["price basis", "basis", "priced per"],
    }
    lookup = {re.sub(r"[^a-z0-9%]+", " ", str(c).lower()).strip(): c for c in columns}
    return {field: next((lookup[x] for x in names if x in lookup), None) for field, names in aliases.items()}


def _term_key(value: Any) -> str:
    label = re.sub(r"[^a-z0-9]+", " ", _text(value).lower()).strip()
    aliases = {
        "quote validity": "quote_validity", "prices": "currency", "price currency": "currency",
        "delivery": "delivery_location", "delivery location": "delivery_location",
        "payment": "payment_terms", "payment term": "payment_terms", "payment terms": "payment_terms",
        "tax": "taxes", "gst": "taxes", "freight": "freight",
        "delivery requirement": "delivery_requirement", "price basis": "price_basis",
        "partial quotation": "partial_quotations", "partial quotations": "partial_quotations",
        "award split": "award_split", "split award": "award_split",
    }
    return aliases.get(label, re.sub(r"[^a-z0-9]+", "_", label).strip("_"))


def _spreadsheet(name: str, data: bytes) -> dict | None:
    if Path(name).suffix.lower() not in {".xlsx", ".xls", ".csv"}:
        return None
    sheets = {"CSV": pd.read_csv(io.BytesIO(data), header=None)} if name.lower().endswith(".csv") else pd.read_excel(io.BytesIO(data), sheet_name=None, header=None)
    quotes, extra, commercial_terms = [], [], {}
    quote_table_found = False
    for sheet_name, frame in sheets.items():
        if frame.empty:
            continue
        header_row = None
        header_values = None
        for row_index in range(min(len(frame), 15)):
            values = frame.iloc[row_index].fillna("").astype(str).tolist()
            headers = _map_headers(values)
            keys = {re.sub(r"[^a-z0-9]+", " ", value.lower()).strip() for value in values}
            has_terms = bool(keys & {"term", "commercial term", "term name"}) and bool(
                keys & {"vendor response", "supplier response", "response", "answer", "value"})
            if headers["item"] or has_terms:
                header_row, header_values = row_index, values
                quote_table_found = quote_table_found or bool(headers["item"])
                break
        if header_row is None:
            extra.append(f"Could not recognize quote or terms table headers in sheet '{sheet_name}'.")
            continue
        frame = frame.iloc[header_row + 1:].copy()
        frame.columns = header_values
        frame = frame.reset_index(drop=True)
        mapping = _map_headers(frame.columns)
        if not mapping["item"]:
            header_keys = {re.sub(r"[^a-z0-9]+", " ", str(column).lower()).strip(): column for column in frame.columns}
            term_col = next((header_keys[x] for x in ("term", "commercial term", "term name") if x in header_keys), None)
            value_col = next((header_keys[x] for x in ("vendor response", "supplier response", "response", "answer", "value") if x in header_keys), None)
            if term_col is not None and value_col is not None:
                for _, row in frame.iterrows():
                    term = _term_key(row[term_col])
                    response = _text(row[value_col])
                    if term and response:
                        commercial_terms[term] = response
            else:
                extra.append(f"Could not map response columns in sheet '{sheet_name}'.")
            continue
        for row_idx, row in frame.iterrows():
            item_name = _text(row[mapping["item"]])
            if not item_name:
                continue
            quote = {field: _text(row[col]) if col else "" for field, col in mapping.items()}
            quote["item"] = item_name
            quote["quantity"] = _number(quote["quantity"])
            quote["unit_price"] = _number(quote["unit_price"])
            quote["currency"] = quote["currency"].upper()
            quote["price_basis"] = quote["price_basis"] or f"per {quote['unit']}"
            quote.update({"confidence": "high", "evidence": "; ".join(f"{field}: {quote[field]}" for field in ("item", "quantity", "unit", "unit_price", "currency") if quote.get(field)),
                          "source_location": f"{sheet_name}, row {int(row_idx) + header_row + 2}", "notes": ""})
            quotes.append(quote)
    notes = list(extra)
    if not quote_table_found:
        notes.insert(0, "Quote table format not recognized; line-item coverage cannot be determined from this workbook.")
    return {"file_name": name, "vendor": Path(name).stem,
            "source_status": "readable" if quote_table_found else "unrecognized_layout",
            "items": quotes, "questionnaire": {}, "commercial_terms": commercial_terms,
            "source_notes": "\n".join(notes)}


def _prepare_unstructured(name: str, data: bytes) -> dict:
    ext = Path(name).suffix.lower()
    if ext in {".docx"}:
        attachments = []
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            for member in archive.namelist():
                if member.startswith("word/media/") and not member.endswith("/"):
                    mime = mimetypes.guess_type(member)[0]
                    if mime and mime.startswith("image/"):
                        attachments.append({"name": Path(member).name, "bytes": archive.read(member), "mime": mime})
        return {"name": name, "text": _docx_text(data), "attachments": attachments}
    if ext == ".pdf":
        text = _pdf_text(data)
        try:
            from pypdf import PdfReader
            has_images = any(bool(page.images) for page in PdfReader(io.BytesIO(data)).pages)
        except Exception:
            has_images = False
        return {"name": name, "bytes": data, "mime": "application/pdf"} if has_images or len(text) < 20 else {"name": name, "text": text}
    if ext == ".txt":
        return {"name": name, "text": data.decode("utf-8", errors="replace")}
    if ext == ".eml":
        message = BytesParser(policy=policy.default).parsebytes(data)
        body = []
        if message.is_multipart():
            for part in message.walk():
                if part.get_content_type() in {"text/plain", "text/html"} and not part.get_filename():
                    text_part = str(part.get_content())
                    body.append(re.sub(r"\s+", " ", unescape(re.sub(r"<[^>]+>", " ", text_part))) if part.get_content_type() == "text/html" else text_part)
        elif message.get_content_type() == "text/plain":
            body.append(message.get_content())
        elif message.get_content_type() == "text/html":
            body.append(re.sub(r"\s+", " ", unescape(re.sub(r"<[^>]+>", " ", str(message.get_content())))))
        subject = _text(message.get("subject"))
        content = "\n".join(map(str, body)).strip()
        if not content:
            raise ValueError(f"{name} has no readable email body. Attach the quote separately or upload a readable message.")
        return {"name": name, "text": f"Email subject: {subject}\n\n{content}"}
    if ext in {".png", ".jpg", ".jpeg", ".webp"}:
        mime = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}[ext]
        try:
            from PIL import Image, ImageOps
            image = ImageOps.exif_transpose(Image.open(io.BytesIO(data)))
            if max(image.size) > 3000:
                image.thumbnail((3000, 3000))
            out = io.BytesIO()
            if mime == "image/png":
                image.save(out, format="PNG", optimize=True)
            else:
                image.convert("RGB").save(out, format="JPEG", quality=90, optimize=True)
                mime = "image/jpeg"
            data = out.getvalue()
        except Exception:
            pass
        return {"name": name, "bytes": data, "mime": mime}
    raise ValueError(f"Unsupported file type: {ext or '(no extension)'}")


def _extract_batch(documents: list[dict], rfx: dict) -> list[dict]:
    from google.genai import types
    has_buyer_rfx = bool(rfx.get("items"))
    reference = {"scope": rfx.get("scope", ""), "items": rfx.get("items", []),
                 "questions": rfx.get("questionnaire", []), "terms": rfx.get("commercial_terms", {})}
    questionnaire_rule = ("Return one questionnaire array entry for every buyer question, in the same order, with the exact question in 'question' and the supplier's wording in 'answer'. Use an empty answer if the vendor did not answer it."
                          if reference["questions"] else
                          "Extract questionnaire content as an array of objects with 'question' and 'answer' strings; return an empty array if none are present.")
    terms_rule = ("Return commercial terms as an array of objects with the exact buyer term key in 'term' and supplier wording in 'response'. Include one entry per buyer term and use an empty response if unstated."
                  if reference["terms"] else
                  "Extract commercial terms as an array of objects with 'term' and 'response' strings; return an empty array if none are present.")
    line_rule = ("Set rfx_line_id to the 1-based matching RFx line, or an empty string if uncertain."
                 if has_buyer_rfx else "Set rfx_line_id to an empty string because no buyer RFx line list was provided.")
    prompt = f"""Read each labeled supplier response. Return JSON only with a top-level vendors array containing one object per source file. Each object must have file_name, vendor, source_status (readable/partially_readable/unreadable), source_notes, items (array of rows), questionnaire (array of question/answer objects), and commercial_terms (array of term/response objects). To keep the response compact, encode each item row as exactly 14 strings in this order: [rfx_line_id, item, specification, quantity, unit, unit_price, currency, price_basis, confidence, evidence, source_location, gst, delivery_days, notes]. Use an empty string for an unknown value; use a short exact evidence phrase. {line_rule} {questionnaire_rule} {terms_rule} Extract only supported facts; do not fill gaps by guessing. Preserve supplier wording in evidence. For parts you cannot read, state that explicitly and use low confidence/empty values rather than calling the information vendor-missing. Include file_name exactly as labeled.
Buyer RFx, if supplied: {json.dumps(reference, ensure_ascii=False)}"""
    contents: list[Any] = [prompt]
    for doc in documents:
        contents.append(f"\nSOURCE FILE: {doc['name']}\n")
        if "text" in doc:
            if len(doc["text"]) > MAX_SOURCE_TEXT_CHARS:
                raise ValueError(f"{doc['name']} is longer than the {MAX_SOURCE_TEXT_CHARS:,}-character safe batch limit. Split or simplify this response; no text was silently dropped.")
            contents.append(doc["text"])
        for attachment in doc.get("attachments", []):
            contents.append(f"\nEmbedded image in {doc['name']}: {attachment['name']}\n")
            contents.append(types.Part.from_bytes(data=attachment["bytes"], mime_type=attachment["mime"]))
        if "bytes" in doc:
            contents.append(types.Part.from_bytes(data=doc["bytes"], mime_type=doc["mime"]))
    client = _client()
    response = client.models.generate_content(
        model=MODEL, contents=contents,
        config=types.GenerateContentConfig(response_mime_type="application/json", response_schema=VENDOR_RESPONSE_SCHEMA,
                                           max_output_tokens=MAX_OUTPUT_TOKENS),
    )
    try:
        parsed = _json(response.text)
    except (TypeError, ValueError) as exc:
        reason = str(exc).splitlines()[0][:240] or type(exc).__name__
        status = _gemini_finish_status(response)
        raise ValueError(f"Gemini response could not be parsed as JSON ({type(exc).__name__}: {reason}; {status}). "
                         "No comparison was created. Your uploaded files are still available. "
                         "This request was not retried automatically.") from exc
    vendors = parsed.get("vendors", [])
    if not isinstance(vendors, list):
        raise ValueError("Gemini response is missing the vendors list.")
    return vendors


def _normalise(raw: dict, fallback_name: str, rfx: dict | None = None) -> dict:
    rfx = rfx or empty_rfx()
    items = []
    item_fields = ("rfx_line_id", "item", "specification", "quantity", "unit", "unit_price", "currency",
                   "price_basis", "confidence", "evidence", "source_location", "gst", "delivery_days", "notes")
    for q in raw.get("items", []) or []:
        if isinstance(q, (list, tuple)):
            q = dict(zip(item_fields, q))
        if not isinstance(q, dict):
            continue
        price = _number(q.get("unit_price"))
        currency = _text(q.get("currency")).upper()
        currency = {"₹": "INR", "RS": "INR", "RUPEE": "INR", "RUPEES": "INR",
                    "US DOLLAR": "USD", "US DOLLARS": "USD"}.get(currency, currency)
        notes = _text(q.get("notes"))
        assumed_currency = not currency
        if assumed_currency:
            notes = (notes + "; Currency not stated by vendor.").strip("; ")
        confidence = (_text(q.get("confidence")) or "medium").lower()
        confidence = confidence if confidence in {"high", "medium", "low"} else "medium"
        items.append({
            "item": _text(q.get("item")), "specification": _text(q.get("specification")),
            "rfx_line_id": _number(q.get("rfx_line_id")), "quantity": _number(q.get("quantity")),
            "unit": _text(q.get("unit")), "unit_price": price, "currency": currency,
            "currency_assumed": assumed_currency, "gst": _text(q.get("gst")),
            "delivery_days": _text(q.get("delivery_days")),
            "price_basis": _text(q.get("price_basis")),
            "confidence": confidence,
            "evidence": _text(q.get("evidence")), "source_location": _text(q.get("source_location")), "notes": notes,
        })
    qdata = raw.get("questionnaire") or {}
    questions = rfx.get("questionnaire", [])
    if isinstance(qdata, list):
        if qdata and isinstance(qdata[0], dict):
            question_aliases = {re.sub(r"\s+", " ", str(question)).strip().casefold(): str(question) for question in questions}
            qdata = {
                question_aliases.get(re.sub(r"\s+", " ", _text(entry.get("question"))).strip().casefold(), _text(entry.get("question"))): _text(entry.get("answer"))
                for entry in qdata if isinstance(entry, dict) and _text(entry.get("question"))
            }
        else:
            qdata = {str(question): _text(ans) for question, ans in zip(questions, qdata)}
    if isinstance(qdata, dict) and questions:
        question_aliases = {re.sub(r"\s+", " ", str(question)).strip().casefold(): str(question) for question in questions}
        qdata = {
            question_aliases.get(re.sub(r"\s+", " ", str(key)).strip().casefold(), str(key)): value
            for key, value in qdata.items()
        }
    if not rfx.get("questionnaire") and isinstance(qdata, dict):
        questionnaire = {str(key): _text(value) for key, value in qdata.items()}
    else:
        questionnaire = {str(question): _text(qdata.get(str(question), qdata.get(f"question_{index + 1}", "")))
                         for index, question in enumerate(questions)}
    status = _text(raw.get("source_status")).lower().replace(" ", "_").replace("-", "_")
    if status == "unrecognized_layout" or "unrecognized_layout" in status:
        status = "unrecognized_layout"
    elif "unreadable" in status:
        status = "unreadable"
    elif "partial" in status:
        status = "partially_readable"
    else:
        status = "readable"
    raw_terms = raw.get("commercial_terms") or {}
    term_aliases = {_key(key): key for key in (rfx.get("commercial_terms") or {})}
    if isinstance(raw_terms, list):
        terms = {}
        for entry in raw_terms:
            if isinstance(entry, dict):
                term, value = _text(entry.get("term")), entry.get("response", "")
            elif isinstance(entry, (list, tuple)) and len(entry) >= 2:
                term, value = _text(entry[0]), entry[1]
            else:
                continue
            if term:
                terms[term_aliases.get(_key(term), term.strip().lower().replace(" ", "_"))] = _text(value)
    elif isinstance(raw_terms, dict):
        terms = {term_aliases.get(_key(key), key): _text(value) for key, value in raw_terms.items()}
    else:
        terms = {}
    return {
        "file_name": fallback_name, "vendor": _text(raw.get("vendor")) or Path(fallback_name).stem,
        "source_status": status,
        "source_notes": _text(raw.get("source_notes")), "items": items,
        "questionnaire": questionnaire,
        "commercial_terms": terms,
    }


def _key(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", (value or "").lower())


def _unit(value: str) -> str:
    key = _key(value)
    return {"pc": "pcs", "piece": "pcs", "pieces": "pcs", "sheet": "sheets", "kilogram": "kg", "kilograms": "kg"}.get(key, key)


def _price_basis_conversion(basis: str, quoted_unit: str) -> tuple[float, bool, str]:
    """Return the count of quoted units per stated price, or explain why unsafe."""
    text = _text(basis).lower().strip()
    unit = _unit(quoted_unit)
    if not text:
        return 1, False, "price basis is not stated"
    batch = re.search(r"\bper\s+(\d+(?:\.\d+)?)\s*([a-z]+)?\b", text)
    if batch:
        count = float(batch.group(1))
        basis_unit = _unit(batch.group(2) or quoted_unit)
        if count > 0 and unit and basis_unit == unit:
            return count, True, ""
        return 1, False, f"price basis '{basis}' cannot be converted to vendor unit '{quoted_unit}'"
    single = re.search(r"\bper\s+([a-z]+)\b", text)
    if single:
        basis_word = _unit(single.group(1))
        if basis_word in {"unit", "item", "each"}:
            return 1, bool(unit), "" if unit else "vendor unit is missing"
        if unit and basis_word == unit:
            return 1, True, ""
        return 1, False, f"price basis '{basis}' cannot be converted to vendor unit '{quoted_unit}'"
    if text in {"each", "unit price", "per unit price"}:
        return 1, bool(unit), "" if unit else "vendor unit is missing"
    return 1, False, f"unrecognized price basis '{basis}'"


def _item_match_score(requested: dict, quote: dict) -> float:
    item_score = SequenceMatcher(None, _key(requested["item"]), _key(quote.get("item", ""))).ratio()
    quoted_spec = _text(quote.get("specification"))
    if quoted_spec and _text(requested.get("specification")):
        spec_score = SequenceMatcher(None, _key(requested["specification"]), _key(quoted_spec)).ratio()
        return 0.8 * item_score + 0.2 * spec_score
    return item_score


def compare_vendor_with_rfx(vendor: dict, rfx: dict | None = None, fx_rates: dict | None = None,
                           base_currency: str | None = None, fx_rate_dates: dict | None = None,
                           fx_rate_source: str | None = None, fx_rate_error: str | None = None) -> dict:
    if not isinstance(rfx, dict):
        raise ValueError("Create an RFx from the buyer's instructions before comparing vendor responses.")
    base_currency = _currency_code(base_currency) or infer_rfx_currency(rfx)
    fx_rates = {k.upper(): float(v) for k, v in (fx_rates or {}).items()}
    fx_rate_dates = {k.upper(): _text(v) for k, v in (fx_rate_dates or {}).items()}
    quotes = vendor.get("items", [])
    matched: set[int] = set()
    lines, missing, qty_issues, unit_issues, currency_issues, spec_issues, other = [], [], [], [], [], [], []
    for line_number, req in enumerate(rfx["items"], 1):
        requested_qty = _number(req.get("quantity"))
        candidates = []
        for idx, quote in enumerate(quotes):
            if idx in matched:
                continue
            line_id = quote.get("rfx_line_id")
            if line_id:
                score = 1.0 if int(line_id) == line_number else -1.0
                target_tied = False
            else:
                target_scores = sorted((_item_match_score(candidate, quote) for candidate in rfx["items"]), reverse=True)
                score = _item_match_score(req, quote)
                # When an item name is duplicated, require the quoted specification
                # to be near the strongest candidate before allowing that RFx row.
                is_best_target = score >= target_scores[0] - 0.025
                target_tied = (is_best_target and target_scores[0] >= 0.65 and len(target_scores) > 1
                               and target_scores[0] - target_scores[1] < 0.025)
                if not is_best_target:
                    score = -1.0
            candidates.append((score, idx, quote, target_tied))
        candidates.sort(reverse=True, key=lambda x: x[0])
        score, idx, quote, target_tied = candidates[0] if candidates else (0, -1, {}, False)
        tied = len(candidates) > 1 and score - candidates[1][0] < 0.06
        confidence = _text(quote.get("confidence")).lower()
        safe_match = score >= 0.82 and not tied and not target_tied and confidence != "low"
        if not safe_match:
            source_status = vendor.get("source_status", "readable")
            if source_status == "unrecognized_layout":
                status = "Needs review - quote table format not recognized"
            elif source_status == "unreadable":
                status = "Unable to read source"
            elif source_status == "partially_readable":
                status = "No match found; source partially unreadable"
            else:
                status = "Needs review" if (score >= 0.65 or target_tied or confidence == "low") else "Not quoted in readable response"
            if status == "Not quoted in readable response":
                missing.append(req["item"])
            if status in {"Needs review", "Unable to read source", "No match found; source partially unreadable"}:
                other.append({"item": req["item"], "issue": status, "evidence": _text(quote.get("evidence"))})
            lines.append({**req, "quote_status": status, "vendor_item": "", "quoted_quantity": None,
                          "quoted_unit": "", "unit_price": None, "currency": "", "normalized_price": None,
                          "comparable_total": None, "confidence": confidence, "evidence": _text(quote.get("evidence")),
                          "source_location": _text(quote.get("source_location")),
                          "notes": _text(quote.get("notes"))})
            continue
        matched.add(idx)
        quoted_qty = _number(quote.get("quantity"))
        if quoted_qty is not None and requested_qty is not None and quoted_qty != requested_qty:
            qty_issues.append({"item": req["item"], "requested": requested_qty, "quoted": quoted_qty})
        requested_unit, quoted_unit = _unit(req["unit"]), _unit(quote.get("unit", ""))
        unit_ok = quoted_unit == requested_unit
        if quote.get("unit") and not unit_ok:
            unit_issues.append({"item": req["item"], "requested": req["unit"], "quoted": quote["unit"]})
        currency = _text(quote.get("currency")).upper() or "Unknown"
        fx_rate = 1.0 if base_currency and currency == base_currency else fx_rates.get(currency)
        if currency == "Unknown":
            currency_issues.append({"item": req["item"], "requested": f"Currency stated for {base_currency or 'comparison'}", "quoted": "Not stated", "issue": "Currency is missing; no conversion or comparable total calculated"})
        elif not base_currency:
            currency_issues.append({"item": req["item"], "requested": "Comparison currency in RFx", "quoted": currency,
                                   "issue": "RFx comparison currency is not specified; price is not normalized"})
        elif currency != base_currency and fx_rate is None:
            currency_issues.append({"item": req["item"], "requested": base_currency, "quoted": currency,
                                   "issue": f"No exchange rate available from {currency} to {base_currency}; price is not normalized"})
        basis = _text(quote.get("price_basis"))
        basis_count, basis_ok, basis_reason = _price_basis_conversion(basis, quote.get("unit", ""))
        unit_price = quote.get("unit_price")
        normalized_price = unit_price / basis_count * fx_rate if unit_price is not None and fx_rate is not None and basis_ok and unit_ok else None
        price_reasons = []
        if not _text(quote.get("unit")):
            price_reasons.append("vendor unit is missing")
        elif not unit_ok:
            price_reasons.append(f"vendor unit '{quote.get('unit')}' differs from RFx unit '{req.get('unit', '')}'")
        if not basis_ok:
            price_reasons.append(basis_reason)
        if unit_price is None:
            price_reasons.append("unit price is missing or unreadable")
        if currency == "Unknown":
            price_reasons.append("vendor did not state a currency")
        elif not base_currency:
            price_reasons.append("RFx comparison currency is not specified")
        elif currency != base_currency and fx_rate is None:
            price_reasons.append(f"no exchange rate is available from {currency} to {base_currency}")
        if normalized_price is None and not price_reasons:
            price_reasons.append("price details could not be safely normalized")
        enough_qty = requested_qty is None or (quoted_qty is not None and quoted_qty >= requested_qty)
        total = normalized_price * requested_qty if requested_qty is not None and normalized_price is not None and unit_ok and enough_qty else None
        if unit_price is not None and not unit_ok:
            other.append({"item": req["item"], "issue": "Unit price is not comparable until units are reconciled"})
        if unit_price is not None and not basis_ok:
            other.append({"item": req["item"], "issue": basis_reason, "evidence": _text(quote.get("evidence"))})
        if requested_qty is not None and quoted_qty is not None and quoted_qty < requested_qty:
            other.append({"item": req["item"], "issue": "Quoted quantity is below the RFx requirement; full-line total omitted"})
        if unit_price is None:
            other.append({"item": req["item"], "issue": "Unit price is missing or unreadable", "evidence": _text(quote.get("evidence"))})
        if quoted_qty is None:
            other.append({"item": req["item"], "issue": "Quoted quantity is missing or unreadable", "evidence": _text(quote.get("evidence"))})
        if not _text(quote.get("unit")):
            other.append({"item": req["item"], "issue": "Quoted unit is missing or unreadable", "evidence": _text(quote.get("evidence"))})
        if not _text(quote.get("delivery_days")) and not _text(vendor.get("commercial_terms", {}).get("delivery_requirement")):
            other.append({"item": req["item"], "issue": "Delivery time is missing or unreadable", "evidence": _text(quote.get("evidence"))})
        if not _text(quote.get("gst")) and not _text(vendor.get("commercial_terms", {}).get("taxes")):
            other.append({"item": req["item"], "issue": "GST/tax treatment is missing or unreadable", "evidence": _text(quote.get("evidence"))})
        if quote.get("currency_assumed"):
            other.append({"item": req["item"], "issue": "Vendor did not state a currency", "evidence": _text(quote.get("evidence"))})
        quoted_spec = _text(quote.get("specification"))
        if quoted_spec and req.get("specification"):
            req_dims = re.findall(r"\d+", req["specification"])
            quoted_dims = re.findall(r"\d+", quoted_spec)
            if req_dims and quoted_dims and req_dims[:3] != quoted_dims[:3]:
                spec_issues.append({"item": req["item"], "requested": req["specification"], "quoted": quoted_spec})
        if _text(quote.get("notes")):
            other.append({"item": req["item"], "issue": _text(quote["notes"]), "evidence": _text(quote.get("evidence"))})
        if confidence in {"low", "medium"}:
            other.append({"item": req["item"], "issue": f"Extraction confidence: {confidence}", "evidence": _text(quote.get("evidence"))})
        lines.append({**req, "quantity": requested_qty, "quote_status": "Quoted", "vendor_item": quote.get("item", ""),
                      "quoted_quantity": quoted_qty, "quoted_unit": quote.get("unit", ""),
                      "unit_price": unit_price, "currency": currency, "price_basis": basis,
                      "price_not_comparable_reason": "; ".join(price_reasons) if normalized_price is None else "",
                      "normalized_price": normalized_price, "fx_rate": fx_rate,
                      "fx_rate_date": fx_rate_dates.get(currency, ""), "comparison_currency": base_currency,
                      "comparable_total": total, "gst": quote.get("gst", ""),
                      "delivery_days": quote.get("delivery_days", ""), "confidence": confidence,
                      "evidence": _text(quote.get("evidence")), "source_location": _text(quote.get("source_location")),
                      "notes": _text(quote.get("notes"))})
    for idx, quote in enumerate(quotes):
        if idx not in matched:
            other.append({"item": _text(quote.get("item")) or "Unmatched quote", "issue": "Could not safely match to an RFx line",
                          "evidence": _text(quote.get("evidence"))})
    questionnaire = vendor.get("questionnaire", {})
    unknown_answers = {"not specified", "unknown", "n/a", "na", "none", "not provided", "unreadable", "-"}
    questions = rfx.get("questionnaire", [])
    questionnaire_issues = [{"question": question, "issue": "No usable answer extracted", "answer": _text(questionnaire.get(question))}
                            for question in questions
                            if not _text(questionnaire.get(question)) or _text(questionnaire.get(question)).lower().strip(" .") in unknown_answers]
    expected_terms = rfx.get("commercial_terms", {})
    actual_terms = vendor.get("commercial_terms", {})
    term_issues = []
    for key, expected in expected_terms.items():
        actual = _text(actual_terms.get(key))
        if not actual:
            term_issues.append({"term": key, "expected": expected, "actual": "Not stated"})
            continue
        if key == "currency" and "unless otherwise stated" in _text(expected).lower():
            continue
        if key == "delivery_requirement":
            expected_days, actual_days = _number(expected), _number(actual)
            if expected_days is not None and actual_days is not None and actual_days <= expected_days:
                continue
        if key == "taxes" and "gst" in actual.lower() and re.search(r"extra|separat", actual, re.I):
            continue
        if _key(_text(expected)) not in _key(actual) and _key(actual) not in _key(_text(expected)):
            term_issues.append({"term": key, "expected": expected, "actual": actual, "issue": "Review deviation from RFx"})
    fx_rates_used = {currency: {"rate": rate, "date": fx_rate_dates.get(currency, "")}
                     for currency, rate in fx_rates.items() if currency != base_currency}
    return {"items": lines, "comparison_currency": base_currency, "fx_rates_used": fx_rates_used,
            "fx_rate_dates": fx_rate_dates, "fx_rate_source": fx_rate_source,
            "fx_rate_error": fx_rate_error, "missing_items": missing, "quantity_mismatches": qty_issues,
            "unit_mismatches": unit_issues, "specification_mismatches": spec_issues,
            "currency_mismatches": currency_issues, "questionnaire_issues": questionnaire_issues,
            "term_issues": term_issues, "other_issues": other, "vendor_notes": vendor.get("source_notes", "")}


def process_vendor_responses(sources, rfx: dict | None = None, fx_rates: dict | None = None,
                             comparison_currency: str | None = None) -> tuple[list[dict], int]:
    """Parse spreadsheets locally, extract all other uploads in one request, then compare locally."""
    if not isinstance(rfx, dict) or not rfx.get("items"):
        raise ValueError("Generate an RFx with at least one line item before processing vendor responses.")
    base_currency = _currency_code(comparison_currency) or infer_rfx_currency(rfx)
    if not base_currency:
        raise ValueError("Set a comparison currency in the RFx commercial terms or enter its three-letter code.")
    structured, unstructured = [], []
    for source in sources:
        name, data = _source(source)
        local = _spreadsheet(name, data)
        if local is not None:
            structured.append(_normalise(local, name, rfx))
        else:
            unstructured.append(_prepare_unstructured(name, data))
    extraction_calls = 0
    if unstructured:
        extraction_calls = 1
        try:
            raw_vendors = _extract_batch(unstructured, rfx)
            if not isinstance(raw_vendors, list):
                raise ValueError("Gemini response did not contain a vendor list.")
            expected = [doc["name"].lower() for doc in unstructured]
            returned = [Path(_text(v.get("file_name"))).name.lower()
                        for v in raw_vendors if isinstance(v, dict)]
            if len(raw_vendors) != len(unstructured) or sorted(returned) != sorted(expected):
                missing = sorted(set(expected) - set(returned))
                extra = sorted(set(returned) - set(expected))
                detail = []
                if missing:
                    detail.append("missing: " + ", ".join(missing))
                if extra:
                    detail.append("unexpected: " + ", ".join(extra))
                if not detail:
                    detail.append("duplicate or malformed vendor entries")
                raise ValueError("Gemini returned incomplete vendor extraction data (" + "; ".join(detail) + ").")
            by_name = {Path(_text(v.get("file_name"))).name.lower(): v for v in raw_vendors}
            structured.extend(_normalise(by_name[doc["name"].lower()], doc["name"], rfx)
                              for doc in unstructured)
        except Exception as exc:
            raise ValueError(f"Vendor extraction failed after {extraction_calls} Gemini request(s) in this attempt. "
                             f"No comparison was created. {exc}") from exc
    source_currencies = [quote.get("currency", "") for vendor in structured for quote in vendor.get("items", [])]
    foreign_currencies = {code for value in source_currencies if (code := _currency_code(value)) and code != base_currency}
    fx_rate_dates: dict[str, str] = {}
    fx_rate_error = None
    fx_rate_source = None
    if fx_rates is None:
        fx_rates, fx_rate_dates, fx_rate_error = _fetch_fx_rates(base_currency, sorted(foreign_currencies))
        fx_rate_source = "Frankfurter" if foreign_currencies else None
    else:
        fx_rate_source = "Provided rates" if foreign_currencies else None
    results = []
    for vendor in structured:
        results.append({"vendor": vendor["vendor"], "file_name": vendor["file_name"], "data": vendor,
                        "terms": {"questionnaire": vendor["questionnaire"], "commercial_terms": vendor["commercial_terms"]},
                        "comparison": compare_vendor_with_rfx(vendor, rfx, fx_rates, base_currency,
                                                                fx_rate_dates, fx_rate_source, fx_rate_error),
                        "status": "Processed"})
    return results, extraction_calls


def process_vendor_response(source, rfx: dict | None = None, fx_rates: dict | None = None) -> dict:
    """Backward-compatible single-file entry point."""
    results, calls = process_vendor_responses([source], rfx, fx_rates)
    results[0]["model_calls"] = calls
    return results[0]


def draft_rfx(description: str) -> dict:
    if not _text(description):
        raise ValueError("Enter the buyer's sourcing instructions before generating an RFx.")
    prompt = f"""Create an RFx using only the buyer instructions below. Return JSON with scope, items (item_id, item, specification, quantity, unit), questionnaire (list), and commercial_terms (object). Preserve all item IDs, item names, specifications, quantities, units, questionnaire content, and terms explicitly given by the buyer. Do not add products, quantities, dates, locations, certifications, terms, or other procurement facts from memory or from a template. Do not assume values not stated; use null for unknown item quantities and omit unknown terms. If essential details are missing, add concise questions to the questionnaire rather than inventing answers.\nBuyer instructions: {description}"""
    from google.genai import types
    client = _client()
    response = client.models.generate_content(model=MODEL, contents=prompt,
        config=types.GenerateContentConfig(response_mime_type="application/json", max_output_tokens=5000))
    result = _json(response.text)
    result["scope"] = _text(result.get("scope")) or description
    result.setdefault("items", [])
    result.setdefault("questionnaire", [])
    result.setdefault("commercial_terms", {})
    if not isinstance(result["items"], list):
        raise ValueError("RFx draft did not contain an item list. Please refine the buyer brief and try again.")
    result["items"] = [{"item_id": _text(item.get("item_id")) or index, "item": _text(item.get("item")),
                        "specification": _text(item.get("specification")),
                        "quantity": _number(item.get("quantity")), "unit": _text(item.get("unit"))}
                       for index, item in enumerate(result["items"], 1) if isinstance(item, dict)]
    if isinstance(result["questionnaire"], dict):
        result["questionnaire"] = list(result["questionnaire"].values())
    result["questionnaire"] = [_text(q) for q in result["questionnaire"] if _text(q)]
    if not isinstance(result["commercial_terms"], dict):
        result["commercial_terms"] = {}
    return result


def analyst_answer(question: str, results: list[dict], rfx: dict | None = None) -> dict:
    """Answer one buyer question using only the compact, current comparison snapshot."""
    if not isinstance(rfx, dict):
        raise ValueError("Create an RFx and process vendor responses before asking the Analyst.")
    compact = []
    for result in results:
        comp = result["comparison"]
        quote_status_counts = {}
        for line in comp["items"]:
            status = line.get("quote_status") or "Unknown"
            quote_status_counts[status] = quote_status_counts.get(status, 0) + 1
        compact.append({"vendor": result["vendor"], "comparison_currency": comp.get("comparison_currency"),
            "fx_rates_used": comp.get("fx_rates_used", {}), "items": [
            {k: line.get(k) for k in ("item_id", "item", "quantity", "unit", "quote_status", "quoted_quantity",
              "quoted_unit", "unit_price", "currency", "price_basis", "normalized_price", "price_not_comparable_reason", "comparable_total",
              "comparison_currency", "fx_rate", "fx_rate_date", "confidence", "evidence", "source_location", "notes")} for line in comp["items"]],
            "quote_status_counts": quote_status_counts,
            "questionnaire": result["terms"]["questionnaire"], "commercial_terms": result["terms"]["commercial_terms"],
            "missing_items": comp["missing_items"], "quantity_mismatches": comp["quantity_mismatches"],
            "unit_mismatches": comp["unit_mismatches"], "specification_mismatches": comp["specification_mismatches"],
            "currency_mismatches": comp["currency_mismatches"], "questionnaire_issues": comp["questionnaire_issues"],
            "term_issues": comp["term_issues"], "issues": comp["other_issues"], "source_notes": comp["vendor_notes"]})
    prompt = f"""Answer only the buyer's question with concise, evidence-based analysis using the supplied comparison. Do not invent values or compare unlike currencies/units. Each vendor's normalized_price and comparable_total are in that vendor comparison's comparison_currency; unit_price and currency preserve the supplier's original quote. Compare normalized values only when present. If they are absent, explain the stated price_not_comparable_reason. Rate and rate date are included when conversion was applied. For line-item counts, treat each vendor's quote_status_counts as authoritative: count only the exact status 'Quoted' as quoted, and report 'Needs review' and 'Not quoted in readable response' separately. A quantity mismatch does not change a line's quote_status. Never claim a vendor is the only one to quote all lines unless the counts show that every other vendor quoted fewer RFx lines. A 'best quote' is not automatically a defensible award: distinguish price, coverage, quantity/specification issues, questionnaire answers, and commercial terms; state the basis for any comparison, and do not imply a definitive award where no weighting criteria were supplied. For missing-data claims, inspect the actual questionnaire and commercial_terms values: only empty or null values are missing, and any non-empty value is stated. Never describe questionnaire answers or commercial terms as unpopulated if any values are present. Mention assumptions or uncertainty only when relevant to the buyer's question; do not add generic caveats or unrelated categories. Return JSON: {{"answer":"markdown text","table":{{"title":"","columns":[],"rows":[]}} or null,"chart":{{"title":"","labels":[],"values":[]}} or null}}. Table/chart are optional; include only when useful and derive all values from the data.\nQuestion: {question}\nRFx: {json.dumps(rfx, ensure_ascii=False)}\nComparison: {json.dumps(compact, ensure_ascii=False)}"""
    from google.genai import types
    client = _client()
    response = client.models.generate_content(model=MODEL, contents=prompt,
        config=types.GenerateContentConfig(response_mime_type="application/json", max_output_tokens=2500))
    answer = _json(response.text)
    answer.setdefault("answer", "")
    answer.setdefault("table", None)
    answer.setdefault("chart", None)
    return answer
